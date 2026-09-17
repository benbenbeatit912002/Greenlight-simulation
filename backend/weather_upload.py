"""Bounded, read-only validation of the public weather workbook contract.

No uploaded bytes are saved, no formulas evaluated and no model is imported here.
Timestamps in Excel date cells are explicitly UTC under schema weather-xlsx-v1.
"""

from __future__ import annotations

import hashlib
import io
import math
from datetime import datetime, timezone
from pathlib import PurePosixPath
from zipfile import ZipFile

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 10_000
MAX_XML_BYTES = 16 * 1024 * 1024
MAX_ISSUES = 30
SCHEMA_VERSION = "weather-xlsx-v1"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REQUIRED_COLUMNS = (
    "timestamp_utc",
    "outside_temperature_C",
    "outside_relative_humidity_pct",
    "global_radiation_W_m2",
    "wind_speed_m_s",
)
OPTIONAL_COLUMNS = ("sky_temperature_C", "soil_temperature_C", "outside_co2_ppm")
COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
AMSTERDAM_COLUMNS = (
    "time",
    "global radiation",
    "wind speed",
    "air temperature",
    "sky temperature",
    "??",
    "CO2 concentration",
    "day number",
    "RH",
)
AMSTERDAM_BOUNDS = {
    "time": (0, 366 * 86400),
    "global radiation": (0, 1600),
    "wind speed": (0, 100),
    "air temperature": (-80, 65),
    "sky temperature": (-120, 80),
    "CO2 concentration": (100, 5000),
    "day number": (0, 366),
    "RH": (0, 100),
}
# Supported import bounds, not agronomic safety limits or calibration evidence.
BOUNDS = {
    "outside_temperature_C": (-80, 65),
    "outside_relative_humidity_pct": (0, 100),
    "global_radiation_W_m2": (0, 1600),
    "wind_speed_m_s": (0, 100),
    "sky_temperature_C": (-120, 80),
    "soil_temperature_C": (-60, 90),
    "outside_co2_ppm": (100, 5000),
}


class WeatherValidationError(ValueError):
    def __init__(self, issues: list[dict]) -> None:
        self.issues = issues[:MAX_ISSUES]
        super().__init__(self.issues[0]["message"])


def fail(message: str, code: str = "INVALID_WORKBOOK") -> None:
    raise WeatherValidationError([{"code": code, "row": None, "column": None, "message": message}])


def _check_archive(raw: bytes) -> None:
    """Inspect bounded XML before openpyxl, including forged dimensions/formulas."""
    if not raw or len(raw) > MAX_UPLOAD_BYTES:
        fail("Choose a nonempty .xlsx file no larger than 2 MiB.", "FILE_SIZE")
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 256 or len(set(names)) != len(names):
                fail("Workbook contains too many or duplicate ZIP entries.")
            if sum(entry.file_size for entry in entries) > MAX_XML_BYTES:
                fail("Uncompressed workbook exceeds 16 MiB.", "FILE_SIZE")
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                fail("This file is not an Excel .xlsx workbook.")
            cells = 0
            for entry in entries:
                name = entry.filename
                if ".." in PurePosixPath(name).parts or name.startswith("/") or "\\" in name:
                    fail("Unsafe workbook ZIP path.")
                if (
                    entry.flag_bits & 1
                    or name.lower().endswith(".bin")
                    or any(
                        token in name.lower()
                        for token in (
                            "externallinks/",
                            "embeddings/",
                            "activex/",
                            "vbaproject",
                        )
                    )
                ):
                    fail(
                        "Encrypted files, macros, embedded objects and external links are not supported."
                    )
                if not name.endswith((".xml", ".rels")):
                    if entry.is_dir():
                        continue
                    # Relationships can hide worksheet XML in .data/.XML parts.
                    # Reject unsupported package parts rather than skip inspection.
                    fail(
                        "Unsupported workbook part. Use the plain template without images or non-XML attachments."
                    )
                cells = _check_xml_part(archive, entry, cells)
    except WeatherValidationError:
        raise
    except Exception:
        # Do not expose parser internals, local paths or untrusted exception strings.
        fail("The workbook is damaged, unsafe, or not a supported .xlsx file.")


def _timestamp(value) -> datetime:
    if isinstance(value, datetime):
        result = (
            value.replace(tzinfo=timezone.utc)
            if value.tzinfo is None
            else value.astimezone(timezone.utc)
        )
    elif isinstance(value, str):
        try:
            result = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("Use a valid ISO date-time with Z or an explicit UTC offset.") from exc
        if result.tzinfo is None:
            raise ValueError("Text timestamps require Z or an explicit UTC offset.")
        result = result.astimezone(timezone.utc)
    else:
        raise ValueError("Use an Excel date-time cell in UTC or an ISO timestamp with an offset.")
    if not 1900 <= result.year <= 2100 or result.second or result.microsecond:
        raise ValueError("Timestamp must be minute-aligned and between years 1900 and 2100.")
    return result


def _validate_amsterdam(rows, raw: bytes, *, include_records: bool = False) -> dict:
    records, issues = [], []
    blank_seen = False

    def issue(row, column, code, message):
        if len(issues) < MAX_ISSUES:
            issues.append({"row": row, "column": column, "code": code, "message": message})

    for row_number, cells in enumerate(rows, start=2):
        values = [cell.value for cell in cells]
        if all(value is None for value in values):
            blank_seen = True
            continue
        if row_number > MAX_ROWS + 1:
            fail("At most 10,000 weather records are supported.", "ROW_LIMIT")
        if blank_seen:
            issue(
                row_number,
                None,
                "EMPTY_ROW",
                "Remove blank rows between weather records.",
            )
        record = dict(zip(AMSTERDAM_COLUMNS, values))
        for column, value in record.items():
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
            ):
                issue(
                    row_number,
                    column,
                    "NUMERIC",
                    "A finite numeric cell is required. Paste values, not text or formulas.",
                )
            elif column in AMSTERDAM_BOUNDS:
                lo, hi = AMSTERDAM_BOUNDS[column]
                if not lo <= value <= hi:
                    issue(
                        row_number,
                        column,
                        "RANGE",
                        f"Supported input range is {lo:g} to {hi:g}.",
                    )
        records.append(record)
        if len(issues) >= MAX_ISSUES:
            break
    if len(records) < 2:
        issue(None, None, "TOO_SHORT", "Provide at least two weather records.")
    if not issues:
        for index, record in enumerate(records):
            if record["time"] % 300 != 0:
                issue(
                    index + 2,
                    "time",
                    "INTERVAL",
                    "Amsterdam times must align to 300-second boundaries.",
                )
            if index and record["time"] - records[index - 1]["time"] != 300:
                issue(
                    index + 2,
                    "time",
                    "INTERVAL",
                    "Use consecutive 300-second (5-minute) intervals. No duplicates, gaps or automatic filling.",
                )
    if issues:
        raise WeatherValidationError(issues)
    result = {
        "schemaVersion": "amsterdam-weather-xlsx-v2",
        "sha256": hashlib.sha256(raw).hexdigest(),
        "rowCount": len(records),
        "timezone": "not specified in source",
        "intervalMinutes": 5,
        "start": f"{records[0]['time']:.0f} s",
        "end": f"{records[-1]['time']:.0f} s",
        "durationHours": (records[-1]["time"] - records[0]["time"]) / 3600,
        "columns": list(AMSTERDAM_COLUMNS),
        "missingModelColumns": [],
        "simulationReady": False,
        "conversionReady": False,
        "modelReady": False,
        "stored": False,
        "warnings": [
            "Workbook validation passed. Full-model readiness also requires at least 36 hours of coverage and an available GreenLight engine.",
            "The uploaded-weather adapter consumes radiation, wind, air temperature, sky temperature, CO2 concentration and RH; it estimates soil temperature with the upstream soilTempNl function and does not interpret ?? or day number.",
        ],
        "ranges": {
            column: {
                "min": min(row[column] for row in records),
                "max": max(row[column] for row in records),
            }
            for column in AMSTERDAM_COLUMNS
        },
        "preview": records[:5],
    }
    if include_records:
        result["_records"] = records
    return result


def validate_weather_workbook(raw: bytes, *, include_records: bool = False) -> dict:
    _check_archive(raw)
    issues: list[dict] = []

    def issue(row, column, code, message):
        if len(issues) < MAX_ISSUES:
            issues.append({"row": row, "column": column, "code": code, "message": message})

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except Exception:
        fail("The workbook could not be opened. Save it as a standard .xlsx file.")
    try:
        if "Weather" not in workbook.sheetnames or set(workbook.sheetnames) - {
            "Weather",
            "Instructions",
        }:
            fail(
                "Use a Weather sheet, with only an optional Instructions sheet.",
                "SHEET_NAMES",
            )
        sheet = workbook["Weather"]
        if sheet.sheet_state != "visible":
            fail("The Weather sheet must be visible.")
        # Do not trust an uploaded dimension tag (which may claim A1:A1).
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=len(AMSTERDAM_COLUMNS))
        header = [cell.value for cell in next(rows)]
        while header and header[-1] is None:
            header.pop()
        if tuple(header) == AMSTERDAM_COLUMNS:
            return _validate_amsterdam(rows, raw, include_records=include_records)
        if any(not isinstance(value, str) for value in header) or len(set(header)) != len(header):
            fail("Row 1 must contain unique column names without blank gaps.", "HEADERS")
        if set(REQUIRED_COLUMNS) - set(header) or set(header) - set(COLUMNS):
            fail(
                "Use the exact nine Amsterdam template headers in order. The earlier weather-xlsx-v1 core schema is also accepted for existing workbooks.",
                "HEADERS",
            )
        records, times = _read_legacy_rows(rows, header, issues, issue)
        if len(records) < 2:
            issue(None, None, "TOO_SHORT", "Provide at least two weather records.")
        interval = _check_legacy_intervals(records, times, issues, issue)
        if issues:
            raise WeatherValidationError(issues)
        missing = [column for column in OPTIONAL_COLUMNS if records[0].get(column) is None]
        warnings = [
            "Preview only: uploaded weather is not connected to the simulation engine in this release."
        ]
        if missing:
            warnings.append(
                "Additional model weather inputs are missing: "
                + ", ".join(missing)
                + ". No values were inferred."
            )
        else:
            warnings.append(
                "Weather columns are complete, but model compatibility, initialization and look-ahead coverage still require verification."
            )
        result = {
            "schemaVersion": SCHEMA_VERSION,
            "sha256": hashlib.sha256(raw).hexdigest(),
            "rowCount": len(records),
            "timezone": "UTC",
            "intervalMinutes": interval,
            "start": records[0]["timestamp_utc"],
            "end": records[-1]["timestamp_utc"],
            "durationHours": (times[-1] - times[0]).total_seconds() / 3600,
            "columns": header,
            "missingModelColumns": missing,
            "simulationReady": False,
            "conversionReady": False,
            "modelReady": False,
            "stored": False,
            "warnings": warnings,
            "ranges": {
                column: {
                    "min": min(record[column] for record in records),
                    "max": max(record[column] for record in records),
                }
                for column in header
                if column != "timestamp_utc" and records[0][column] is not None
            },
            "preview": records[:5],
        }
        if include_records:
            result["_records"] = records
        return result
    except WeatherValidationError:
        raise
    except Exception:
        fail("The Weather sheet could not be read. Check the template and cell values.")
    finally:
        workbook.close()


def _check_xml_value(element, tag, string_lengths):
    if "macroenabled" in element.attrib.get("ContentType", "").lower():
        fail("Macro-enabled workbooks are not supported.")
    if tag == "f":
        fail(
            "Formulas are not accepted. Paste values only.",
            "FORMULA",
        )
    if tag == "Relationship" and element.attrib.get("TargetMode") == "External":
        fail("External workbook relationships are not supported.")
    if element.text and len(element.text) > 4096:
        fail("Workbook contains an oversized text cell.", "CELL_SIZE")
    if string_lengths and element.text:
        string_lengths[-1] += len(element.text)
        if string_lengths[-1] > 4096:
            fail(
                "Workbook contains an oversized assembled string.",
                "CELL_SIZE",
            )
    if tag in {"si", "is"}:
        length = string_lengths.pop()
        if length > 4096:
            fail(
                "Workbook contains an oversized assembled string.",
                "CELL_SIZE",
            )
        if string_lengths:
            string_lengths[-1] += length


def _check_xml_part(archive, entry, cells):
    with archive.open(entry) as stream:
        worksheet = False
        root_seen = False
        current_row = 0
        previous_row = 0
        seen_cells = set()
        string_lengths = []
        for event, element in iterparse(stream, events=("start", "end")):
            tag = element.tag.rsplit("}", 1)[-1]
            if event == "start":
                if not root_seen:
                    # Identify worksheet content, not its ZIP location.
                    # OOXML relationships may use nonstandard part paths.
                    worksheet = tag == "worksheet"
                    root_seen = True
                if tag in {"si", "is"}:
                    string_lengths.append(0)
                if worksheet and tag == "row":
                    current_row = int(element.attrib.get("r", "0"))
                    if not previous_row < current_row <= MAX_ROWS + 1:
                        fail(
                            "Worksheet row indices must be unique, ordered and within 10,001 rows.",
                            "SHEET_SIZE",
                        )
                    previous_row = current_row
                if worksheet and tag == "c":
                    cells += 1
                    coordinate = element.attrib.get("r", "")
                    row, col = coordinate_to_tuple(coordinate)
                    if (
                        cells > 100_000
                        or not 1 <= row <= MAX_ROWS + 1
                        or not 1 <= col <= len(AMSTERDAM_COLUMNS)
                    ):
                        fail(
                            "Workbook exceeds 10,000 data rows or 9 columns.",
                            "SHEET_SIZE",
                        )
                    if row != current_row:
                        fail(
                            "Cell coordinates do not match their worksheet row.",
                            "CELL_COORDINATE",
                        )
                    if (row, col) in seen_cells:
                        fail(
                            "Duplicate cell coordinates are not accepted.",
                            "DUPLICATE_CELL",
                        )
                    seen_cells.add((row, col))
                continue
            _check_xml_value(element, tag, string_lengths)
            if worksheet and tag == "row":
                current_row = 0
            element.clear()
    return cells


def _check_legacy_numbers(record, header, row_number, issue):
    for column in header:
        if column == "timestamp_utc":
            continue
        value = record[column]
        if value is None and column in OPTIONAL_COLUMNS:
            continue
        minimum, maximum = BOUNDS[column]
        if (
            isinstance(value, bool)
            or not isinstance(value, (float, int))
            or not math.isfinite(value)
        ):
            issue(
                row_number,
                column,
                "NUMERIC",
                "A finite numeric value is required; numeric text is not accepted.",
            )
        elif not minimum <= value <= maximum:
            issue(
                row_number,
                column,
                "RANGE",
                f"Supported input range is {minimum:g} to {maximum:g}.",
            )


def _check_legacy_intervals(records, times, issues, issue):
    if not issues:
        interval = (times[1] - times[0]).total_seconds() / 60
        for index in range(1, len(times)):
            gap = (times[index] - times[index - 1]).total_seconds() / 60
            if gap <= 0:
                issue(
                    index + 2,
                    "timestamp_utc",
                    "ORDER",
                    "Timestamps must be strictly increasing with no duplicates.",
                )
            elif gap != interval or not 1 <= gap <= 60:
                issue(
                    index + 2,
                    "timestamp_utc",
                    "INTERVAL",
                    "Use a regular interval of 1 to 60 minutes. Gaps are not filled automatically.",
                )
        for column in OPTIONAL_COLUMNS:
            count = sum(record.get(column) is not None for record in records)
            if count not in (0, len(records)):
                issue(
                    None,
                    column,
                    "PARTIAL_COLUMN",
                    "Fill every row of an optional column, or leave the entire column blank.",
                )
    return interval if not issues else None


def _read_legacy_rows(rows, header, issues, issue):
    records: list[dict] = []
    times: list[datetime] = []
    blank_row_seen = False
    for row_number, cells in enumerate(rows, start=2):
        values = [cell.value for cell in cells]
        if any(isinstance(value, str) and len(value) > 4096 for value in values):
            fail("Workbook contains an oversized assembled cell value.", "CELL_SIZE")
        if all(value is None for value in values):
            blank_row_seen = True
            continue
        if row_number > MAX_ROWS + 1:
            fail("At most 10,000 weather records are supported.", "ROW_LIMIT")
        if blank_row_seen:
            issue(
                row_number,
                None,
                "EMPTY_ROW",
                "Remove blank rows between weather records.",
            )
        if any(value is not None for value in values[len(header) :]):
            issue(
                row_number,
                None,
                "EXTRA_CELL",
                "Data extends beyond the named columns.",
            )
        record = dict(zip(header, values))
        try:
            moment = _timestamp(record["timestamp_utc"])
            times.append(moment)
            record["timestamp_utc"] = moment.isoformat().replace("+00:00", "Z")
        except (ValueError, TypeError, OverflowError) as exc:
            issue(row_number, "timestamp_utc", "TIMESTAMP", str(exc))
        _check_legacy_numbers(record, header, row_number, issue)
        records.append(record)
        if len(issues) == MAX_ISSUES:
            break
    return records, times
